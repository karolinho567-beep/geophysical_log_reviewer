"""Manifest parsing, TIFF recovery, and row-level audit tests."""
from __future__ import annotations

import csv
import os
import threading
import time
from collections import OrderedDict

import pytest
from openpyxl import Workbook

from logcv.review.imports import (
    ColumnSelectionRequired,
    DepthFilter,
    ImportCancelled,
    audit_report_path,
    load_manifest,
    normalize_identifier,
    read_source,
    resolve_manifest,
    select_columns,
    write_audit_csv,
)


def _touch(path, value=b"test"):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(value)
    return str(path)


def _ok_probe(path):
    return os.path.getsize(path)


def test_headerless_text_preserves_short_and_14_digit_identifiers(tmp_path):
    source = tmp_path / "apis.txt"
    source.write_text("21441\n\n22536\n42175000720000\n", encoding="utf-8-sig")
    table = load_manifest(str(source))
    assert table.identifier_column == "input_value"
    assert table.path_column is None
    assert [row["input_value"] for row in table.rows] == [
        "21441", "22536", "42175000720000",
    ]


def test_csv_and_tsv_aliases_and_quoted_unc_paths_are_detected(tmp_path):
    csv_path = tmp_path / "manifest.csv"
    csv_path.write_text(
        'UWI,TIFPATH,COMMENT\n"21441","\\\\server\\logs\\21441_RES.tif","x"\n',
        encoding="utf-8-sig",
    )
    table = load_manifest(str(csv_path))
    assert (table.identifier_column, table.path_column) == ("UWI", "TIFPATH")
    assert table.rows[0]["TIFPATH"] == r"\\server\logs\21441_RES.tif"

    tsv_path = tmp_path / "manifest.tsv"
    tsv_path.write_text("well_id\timage_path\n22536\tC:\\logs\\22536.tif\n",
                        encoding="utf-8")
    table = load_manifest(str(tsv_path))
    assert (table.identifier_column, table.path_column) == ("well_id", "image_path")


def test_location_and_depth_columns_are_detected_case_insensitively(tmp_path):
    source = tmp_path / "wells.csv"
    source.write_text(
        "API,TIFPATH,LAT,LONG,Depth1,Depth2\n"
        "21441,a.tif,28.8,-96.9,175,925\n",
        encoding="utf-8",
    )
    table = load_manifest(str(source))
    assert table.latitude_column == "LAT"
    assert table.longitude_column == "LONG"
    assert table.depth_column == "Depth1"
    assert table.log_top_depth_column == "Depth1"
    assert table.log_bottom_depth_column == "Depth2"


def test_log_intervals_follow_each_tiff_for_same_identifier(tmp_path):
    first = _touch(tmp_path / "21441_GR.tif")
    second = _touch(tmp_path / "21441_RES.tif")
    source = tmp_path / "intervals.csv"
    source.write_text(
        "API,TIFPATH,DEPTH1,DEPTH2\n"
        f"21441,{first},100,500\n"
        f"21441,{second},450,900\n",
        encoding="utf-8",
    )
    result = resolve_manifest(load_manifest(str(source)), probe_fn=_ok_probe)
    first_key = os.path.normcase(os.path.abspath(first))
    second_key = os.path.normcase(os.path.abspath(second))
    assert result.path_locations[first_key] == {
        "log_top_depth": "100", "log_bottom_depth": "500",
    }
    assert result.path_locations[second_key] == {
        "log_top_depth": "450", "log_bottom_depth": "900",
    }


def test_xlsx_uses_first_sheet_with_identifier_columns_and_formats_numbers(tmp_path):
    path = tmp_path / "manifest.xlsx"
    book = Workbook()
    book.active.title = "notes"
    book.active.append(["description"])
    book.active.append(["not the data"])
    data = book.create_sheet("wells")
    data.append(["API Number", "TIFF Path"])
    data.append([21441, r"C:\logs\21441.tif"])
    book.save(path)
    table = load_manifest(str(path))
    assert table.rows[0]["API Number"] == "21441"
    assert table.identifier_column == "API Number"
    assert table.path_column == "TIFF Path"


def test_ambiguous_columns_require_explicit_mapping(tmp_path):
    path = tmp_path / "ambiguous.csv"
    path.write_text("API,UWI,TIFPATH,FILE_PATH\n1,2,a.tif,b.tif\n", encoding="utf-8")
    with pytest.raises(ColumnSelectionRequired) as caught:
        load_manifest(str(path))
    table = select_columns(caught.value.table, "UWI", "FILE_PATH")
    assert (table.identifier_column, table.path_column) == ("UWI", "FILE_PATH")


def test_identifier_matching_preserves_source_and_adds_only_12_digit_variant():
    assert normalize_identifier("21441") == ("21441", ("21441",))
    assert normalize_identifier("424693343100") == (
        "424693343100", ("424693343100", "42469334310000")
    )
    assert normalize_identifier("42175000720000") == (
        "42175000720000", ("42175000720000",)
    )
    assert normalize_identifier("ABC") == ("", ())


def test_resolution_prefers_folder_then_recovers_external_and_loads_all_tiffs(tmp_path):
    local = tmp_path / "local"
    external = tmp_path / "external"
    local_copy = _touch(local / "21441_RES.tif")
    listed_copy = _touch(external / "21441_RES.tif")
    second = _touch(external / "21441_GR.tif")
    source = tmp_path / "manifest.csv"
    with source.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["UWI", "TIFPATH"])
        writer.writerow(["21441", listed_copy])
        writer.writerow(["21441", second])
        writer.writerow(["21441", listed_copy])
        writer.writerow(["21441", "IHS ASSOC IMAGE"])
    result = resolve_manifest(load_manifest(str(source)), [local_copy], _ok_probe)
    assert result.loaded_identifiers == ["21441"]
    assert result.paths == [os.path.abspath(local_copy), os.path.abspath(second)]
    assert result.stats.loaded_tiffs == 2
    assert result.stats.loaded_from_folder == 1
    assert result.stats.recovered_from_paths == 1
    assert result.stats.duplicate_references == 1
    assert result.stats.placeholder_covered == 1
    statuses = [row.load_status for row in result.audit_rows]
    assert statuses == [
        "loaded_from_selected_folder", "recovered_from_listed_path",
        "loaded_from_selected_folder", "placeholder_covered",
    ]


def test_distinct_candidates_are_probed_concurrently_but_results_keep_source_order(
    tmp_path,
):
    source = tmp_path / "parallel.csv"
    paths = [_touch(tmp_path / f"{value}.tif") for value in range(1, 9)]
    with source.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle)
        writer.writerow(["API", "TIFPATH"])
        for value, path in enumerate(paths, start=1):
            writer.writerow([value, path])
    lock = threading.Lock()
    active = 0
    maximum_active = 0

    def slow_probe(path):
        nonlocal active, maximum_active
        with lock:
            active += 1
            maximum_active = max(maximum_active, active)
        try:
            time.sleep(0.03)
            return os.path.getsize(path)
        finally:
            with lock:
                active -= 1

    progress = []
    result = resolve_manifest(
        load_manifest(str(source)), probe_fn=slow_probe, max_workers=4,
        progress_fn=lambda done, total, name: progress.append((done, total, name)),
    )
    assert maximum_active >= 2
    assert result.paths == [os.path.abspath(path) for path in paths]
    assert progress[-1][:2] == (8, 8)


def test_unavailable_unc_share_is_checked_once_and_paths_are_not_probed(tmp_path):
    source = tmp_path / "network.csv"
    source.write_text(
        "API,TIFPATH\n1,\\\\server\\share\\1.tif\n"
        "2,\\\\server\\share\\2.tif\n",
        encoding="utf-8",
    )
    checked_roots = []
    probed_paths = []
    result = resolve_manifest(
        load_manifest(str(source)),
        probe_fn=lambda path: probed_paths.append(path),
        network_check=lambda root: checked_roots.append(root) or False,
    )
    assert checked_roots == [r"\\server\share"]
    assert probed_paths == []
    assert [row.load_status for row in result.audit_rows] == [
        "network_unavailable_unresolved", "network_unavailable_unresolved",
    ]
    assert result.stats.unavailable_network_shares == 1
    assert result.stats.network_paths_skipped == 2


def test_cancelled_resolution_stops_before_probing(tmp_path):
    tif = _touch(tmp_path / "1.tif")
    source = tmp_path / "cancel.csv"
    source.write_text(f"API,TIFPATH\n1,{tif}\n", encoding="utf-8")
    cancelled = threading.Event()
    cancelled.set()
    with pytest.raises(ImportCancelled):
        resolve_manifest(
            load_manifest(str(source)), probe_fn=lambda _path: pytest.fail(),
            cancel_event=cancelled,
        )


def test_depth_filter_is_strict_and_can_include_or_exclude_blanks(tmp_path):
    paths = {value: _touch(tmp_path / f"{value}.tif") for value in ("1", "2", "3", "4")}
    source = tmp_path / "depths.csv"
    source.write_text(
        "API,TIFPATH,DEPTH1\n"
        f"1,{paths['1']},199.9\n"
        f"2,{paths['2']},200\n"
        f"3,{paths['3']},250\n"
        f"4,{paths['4']},\n",
        encoding="utf-8",
    )
    table = load_manifest(str(source))
    excluded_blanks = resolve_manifest(
        table, probe_fn=_ok_probe,
        depth_filter=DepthFilter("DEPTH1", "less_than", 200, False),
    )
    assert excluded_blanks.loaded_identifiers == ["1"]
    assert [row.load_status for row in excluded_blanks.audit_rows] == [
        "recovered_from_listed_path", "filtered_depth", "filtered_depth",
        "filtered_blank_depth",
    ]
    assert excluded_blanks.stats.depth_filtered_rows == 3
    assert excluded_blanks.stats.depth_filtered_identifiers == 3

    included_blanks = resolve_manifest(
        table, probe_fn=_ok_probe,
        depth_filter=DepthFilter("DEPTH1", "less_than", 200, True),
    )
    assert included_blanks.loaded_identifiers == ["1", "4"]
    assert included_blanks.stats.eligible_unique_identifiers == 2


def test_greater_than_filter_excludes_invalid_depth_and_propagates_coordinates(tmp_path):
    deep = _touch(tmp_path / "21441.tif")
    invalid = _touch(tmp_path / "22536.tif")
    source = tmp_path / "locations.csv"
    source.write_text(
        "UWI,TIFPATH,Latitude,Longitude,well_depth\n"
        f"21441,{deep},28.801,-96.902,201\n"
        f"22536,{invalid},28.9,-97.0,unknown\n",
        encoding="utf-8",
    )
    table = load_manifest(str(source))
    result = resolve_manifest(
        table, probe_fn=_ok_probe,
        depth_filter=DepthFilter("well_depth", "greater_than", 200, False),
    )
    key = os.path.normcase(os.path.abspath(deep))
    assert result.loaded_identifiers == ["21441"]
    assert result.path_locations[key] == {
        "latitude": "28.801", "longitude": "-96.902",
    }
    assert result.audit_rows[1].load_status == "filtered_invalid_depth"
    assert result.stats.invalid_depth_rows == 1


def test_relative_listed_paths_resolve_from_manifest_directory(tmp_path):
    tif = _touch(tmp_path / "images" / "21441.tif")
    source = tmp_path / "manifest.csv"
    source.write_text("UWI,TIFPATH\n21441,images/21441.tif\n", encoding="utf-8")
    result = resolve_manifest(load_manifest(str(source)), probe_fn=_ok_probe)
    assert result.paths == [os.path.abspath(tif)]
    assert result.audit_rows[0].source_tif_path == "images/21441.tif"
    assert result.audit_rows[0].load_status == "recovered_from_listed_path"


def test_resolution_marks_missing_unreadable_and_placeholder_only_ids(tmp_path):
    unreadable = _touch(tmp_path / "bad.tif")
    source = tmp_path / "manifest.csv"
    source.write_text(
        f"UWI,TIFPATH\n1,{tmp_path / 'missing.tif'}\n2,{unreadable}\n"
        "3,IHS ASSOC IMAGE\n4,not-an-image\n",
        encoding="utf-8",
    )

    def probe(path):
        if path == unreadable:
            raise ValueError("broken TIFF")
        return True

    result = resolve_manifest(load_manifest(str(source)), probe_fn=probe)
    assert result.paths == []
    assert [row.load_status for row in result.audit_rows] == [
        "path_missing_unresolved", "unreadable_image_unresolved",
        "placeholder_unresolved", "unsupported_image_path_unresolved",
    ]
    assert result.stats.unrecovered_identifiers == 4


def test_folder_matching_supports_12_digit_ids_without_changing_stored_value(tmp_path):
    tif = _touch(tmp_path / "42469334310000_log.tif")
    source = tmp_path / "ids.txt"
    source.write_text("424693343100\n", encoding="utf-8")
    result = resolve_manifest(load_manifest(str(source)), [tif], _ok_probe)
    assert result.loaded_identifiers == ["424693343100"]
    assert result.path_identifiers[os.path.normcase(os.path.abspath(tif))] == "424693343100"
    assert result.audit_rows[0].load_status == "loaded_from_selected_folder"


def test_conflicting_basename_for_different_identifiers_is_reported(tmp_path):
    first = tmp_path / "one" / "same.tif"
    second = tmp_path / "two" / "same.tif"
    _touch(first)
    _touch(second)
    source = tmp_path / "manifest.csv"
    source.write_text(f"UWI,TIFPATH\n1,{first}\n2,{second}\n", encoding="utf-8")
    result = resolve_manifest(load_manifest(str(source)), probe_fn=_ok_probe)
    assert result.loaded_identifiers == ["1"]
    assert result.audit_rows[1].load_status == "basename_conflict_unresolved"


def test_audit_csv_preserves_source_columns_and_appends_status(tmp_path):
    tif = _touch(tmp_path / "21441.tif")
    source = tmp_path / "manifest.csv"
    source.write_text(f"UWI,TIFPATH,COMMENT\n21441,{tif},source note\n", encoding="utf-8")
    result = resolve_manifest(load_manifest(str(source)), probe_fn=_ok_probe)
    report = audit_report_path(str(source), str(tmp_path / "review.xlsx"))
    write_audit_csv(report, result)
    with open(report, encoding="utf-8-sig", newline="") as handle:
        row = next(csv.DictReader(handle))
    assert row["COMMENT"] == "source note"
    assert row["source_identifier"] == "21441"
    assert row["well_loaded"] == "True"
    assert row["load_status"] == "recovered_from_listed_path"


def test_supplied_raster_summary_shape_is_stable():
    source = os.path.join(os.path.dirname(__file__), "..", "data", "raw",
                          "RasterSummary_464_Model_Wells.CSV")
    if not os.path.exists(source):
        pytest.skip("project fixture not present")
    table = load_manifest(source)
    identifiers = {row[table.identifier_column] for row in table.rows}
    placeholders = sum(
        row[table.path_column].strip().casefold() == "ihs assoc image"
        for row in table.rows
    )
    lengths = [len(value) for value in identifiers]
    assert len(table.rows) == 4081
    assert len(identifiers) == 405
    assert lengths.count(5) == 2
    assert lengths.count(12) == 19
    assert lengths.count(14) == 384
    assert placeholders == 460

    # Exercise candidate construction without touching any listed network path.
    # The injected share check rejects the one UNC root before per-file probing.
    dry_result = resolve_manifest(
        table,
        probe_fn=lambda path: pytest.fail(f"unexpected path probe: {path}"),
        network_check=lambda _root: False,
    )
    assert dry_result.stats.unique_tiff_candidates == 943
    assert dry_result.stats.unavailable_network_shares == 1
    assert dry_result.stats.network_paths_skipped == 952
    assert dry_result.stats.loaded_tiffs == 0
